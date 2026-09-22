"""
Word (.docx) 文件读取器
使用 python-docx 提取段落和表格文本，保留标题层级
"""

import os
import logging
from typing import List

from .base import BaseReader, ReaderResult

logger = logging.getLogger(__name__)


class WordReader(BaseReader):
    """读取 .docx 文件"""

    def can_handle(self, filepath: str) -> bool:
        return filepath.lower().endswith('.docx')

    def read(self, filepath: str) -> ReaderResult:
        if not os.path.exists(filepath):
            return ReaderResult(
                text="", filepath=filepath, file_type="word",
                success=False, error=f"File not found: {filepath}"
            )

        try:
            from docx import Document

            doc = Document(filepath)

            # 提取段落文本，保留标题级别
            paragraphs = []
            heading_count = 0
            para_count = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style_name = para.style.name.lower() if para.style else ""
                if "heading" in style_name or style_name.startswith("h"):
                    # 用 markdown 风格的标题标记
                    level = 1
                    for s in ["heading 1", "heading 2", "heading 3",
                              "heading 4", "heading 5", "heading 6",
                              "h1", "h2", "h3", "h4", "h5", "h6"]:
                        if style_name == s:
                            level = int(s[-1])
                            break
                    paragraphs.append("#" * level + " " + text)
                    heading_count += 1
                else:
                    paragraphs.append(text)
                    para_count += 1

            # 提取表格文本
            table_count = len(doc.tables)
            for table in doc.tables:
                paragraphs.append("\n[TABLE]")
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    paragraphs.append(" | ".join(cells))
                paragraphs.append("[/TABLE]")

            full_text = "\n\n".join(paragraphs)

            return ReaderResult(
                text=full_text,
                filepath=filepath,
                file_type="word",
                metadata={
                    "paragraphs": para_count,
                    "headings": heading_count,
                    "tables": table_count,
                },
                success=True,
            )
        except Exception as e:
            logger.exception("Failed to read Word file: %s", filepath)
            return ReaderResult(
                text="", filepath=filepath, file_type="word",
                success=False, error=str(e)
            )