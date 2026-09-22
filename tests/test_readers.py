"""读取器测试：PDF 文本/扫描分流、OCR 渲染目录清理、编码检测、HTML 类型判定、读取器分发"""

import os
import unittest

import pymupdf
from docx import Document

from file_processor.config import ProcessorConfig
from file_processor.pipeline import FileProcessorPipeline
from file_processor.readers.pdf_reader import PDFReader
from file_processor.readers.txt_reader import TxtReader
from file_processor.readers.html_reader import HTMLReader
from file_processor.readers.word_reader import WordReader
from tests._tempdir import cleanup, workspace_tempdir

TEXT_PDF_SENTENCE = (
    "This document contains a real text layer used for extraction testing. "
    "It is long enough to exceed the minimum character threshold."
)
CHINESE_SENTENCE = "这是一段中文内容，用于验证编码检测是否正确，包含标点符号与数字12345。"


def make_pdf(path, text=""):
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


class StubQwenOCR:
    """替代 Qwen-VL，记录被识别的图片路径"""

    def __init__(self, text="识别出的页面文本内容"):
        self.text = text
        self.calls = []

    def ocr_image(self, image_path):
        self.calls.append(image_path)
        return self.text


class PDFReaderTest(unittest.TestCase):

    def setUp(self):
        self._tmp = workspace_tempdir()

    def tearDown(self):
        cleanup(self._tmp)

    def test_text_pdf_is_read_without_ocr(self):
        path = os.path.join(self._tmp, "text.pdf")
        make_pdf(path, TEXT_PDF_SENTENCE)

        result = PDFReader(qwen_ocr=None).read(path)
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.file_type, "pdf_download")
        self.assertFalse(result.metadata["used_ocr"])
        self.assertIn("real text layer", result.text)
        self.assertEqual(result.metadata["page_count"], 1)

    def test_scanned_pdf_without_api_key_fails_with_clear_message(self):
        path = os.path.join(self._tmp, "scanned.pdf")
        make_pdf(path)  # 无文字层

        result = PDFReader(qwen_ocr=None).read(path)
        self.assertFalse(result.success)
        self.assertIn("API key", result.error)
        self.assertEqual(result.file_type, "pdf_photo")

    def test_scanned_pdf_uses_qwen_and_cleans_work_dir(self):
        path = os.path.join(self._tmp, "scanned.pdf")
        make_pdf(path)
        work_dir = os.path.join(self._tmp, "ocr_temp")
        stub = StubQwenOCR()

        result = PDFReader(qwen_ocr=stub, work_dir=work_dir).read(path)
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.file_type, "pdf_photo")
        self.assertTrue(result.metadata["used_ocr"])
        self.assertIn("识别出的页面文本内容", result.text)
        self.assertEqual(len(stub.calls), 1)
        # 渲染的临时目录必须清理干净
        self.assertEqual(os.listdir(work_dir) if os.path.isdir(work_dir) else [], [])

    def test_missing_file(self):
        result = PDFReader().read(os.path.join(self._tmp, "nope.pdf"))
        self.assertFalse(result.success)
        self.assertIn("not found", result.error)


class TxtReaderTest(unittest.TestCase):

    def setUp(self):
        self._tmp = workspace_tempdir()

    def tearDown(self):
        cleanup(self._tmp)

    def test_utf8_file(self):
        path = os.path.join(self._tmp, "utf8.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write(CHINESE_SENTENCE * 3)

        result = TxtReader().read(path)
        self.assertTrue(result.success, result.error)
        self.assertIn("这是一段中文内容", result.text)
        self.assertEqual(result.file_type, "txt")

    def test_gbk_file_is_decoded_correctly(self):
        path = os.path.join(self._tmp, "gbk.txt")
        with open(path, "wb") as f:
            f.write((CHINESE_SENTENCE * 5).encode("gbk"))

        result = TxtReader().read(path)
        self.assertTrue(result.success, result.error)
        self.assertIn("这是一段中文内容", result.text)
        self.assertNotIn("\ufffd", result.text)

    def test_missing_file(self):
        result = TxtReader().read(os.path.join(self._tmp, "nope.txt"))
        self.assertFalse(result.success)


class HtmlReaderTest(unittest.TestCase):

    def setUp(self):
        self._tmp = workspace_tempdir()

    def tearDown(self):
        cleanup(self._tmp)

    def _write(self, name, content):
        path = os.path.join(self._tmp, name)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def test_source_html_keeps_structure(self):
        path = self._write("source.html", """
        <html><head><title>标题</title></head><body>
          <h1>第一章</h1><h2>1.1 小节</h2>
          <p>这是一个段落。</p><p>这是第二个段落。</p>
          <ul><li>条目一</li><li>条目二</li><li>条目三</li></ul>
          <nav>导航噪声</nav><script>var a=1;</script>
        </body></html>""")
        result = HTMLReader().read(path)
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.file_type, "html_source")
        self.assertIn("# 第一章", result.text)
        self.assertIn("- 条目一", result.text)
        self.assertNotIn("导航噪声", result.text)
        self.assertNotIn("var a=1", result.text)
        self.assertEqual(result.metadata["title"], "标题")

    def test_screenshot_html_is_detected(self):
        images = "".join(
            f'<img src="data:image/png;base64,AAAA{i}">' for i in range(6))
        path = self._write("shot.html", f"""
        <html><body>
          {images}
          <div style="position:absolute;left:10px;top:20px">可见文本内容在此</div>
        </body></html>""")
        result = HTMLReader().read(path)
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.file_type, "html_screenshot")
        self.assertIn("可见文本内容在此", result.text)


class WordReaderTest(unittest.TestCase):

    def setUp(self):
        self._tmp = workspace_tempdir()

    def tearDown(self):
        cleanup(self._tmp)

    def test_docx_text(self):
        path = os.path.join(self._tmp, "doc.docx")
        document = Document()
        document.add_heading("文档标题", level=1)
        document.add_paragraph("这是一个包含中文的段落内容。")
        document.save(path)

        result = WordReader().read(path)
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.file_type, "word")
        self.assertIn("文档标题", result.text)
        self.assertIn("包含中文的段落内容", result.text)


class ReaderDispatchTest(unittest.TestCase):
    """通过 pipeline 验证扩展名 → 读取器 的分发与缓存"""

    def setUp(self):
        self._tmp = workspace_tempdir()
        config = ProcessorConfig()
        config.input_dirs = {"txt": self._tmp}
        config.output_dir = os.path.join(self._tmp, "out")
        config.cache_dir = os.path.join(self._tmp, "cache")
        config.quality_enabled = False
        self.pipeline = FileProcessorPipeline(config)

    def tearDown(self):
        cleanup(self._tmp)

    def test_txt_dispatches_to_txt_reader_and_caches(self):
        path = os.path.join(self._tmp, "a.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write(CHINESE_SENTENCE * 3)

        first = self.pipeline._read_single(path)
        self.assertTrue(first.success)
        self.assertEqual(first.file_type, "txt")

        second = self.pipeline._read_single(path)
        self.assertTrue(second.success)
        self.assertEqual(second.text, first.text)
        self.assertEqual(self.pipeline.stats.cache_hits, 1, "第二次读取应命中缓存")

    def test_unsupported_extension_fails(self):
        path = os.path.join(self._tmp, "data.csv")
        with open(path, "w", encoding="utf-8") as f:
            f.write("a,b,c")
        result = self.pipeline._read_single(path)
        self.assertFalse(result.success)
        self.assertIn("No reader", result.error)

    def test_failed_read_is_not_cached(self):
        path = os.path.join(self._tmp, "missing.txt")
        self.assertFalse(self.pipeline._read_single(path).success)
        self.pipeline._read_single(path)
        self.assertEqual(self.pipeline.stats.cache_hits, 0, "失败结果不应进入缓存")


if __name__ == "__main__":
    unittest.main()